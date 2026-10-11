# -*- coding: utf-8 -*-
"""检视《中国鸟类的生活史和生态学特征数据集》(biodiversity-science 2021, 29(9):1149-1153) xlsx 结构。"""
import sys
import openpyxl

path = sys.argv[1]
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
print("sheets:", wb.sheetnames)
for ws in wb.worksheets:
    print(f"\n=== sheet: {ws.title}  dims={ws.calculate_dimension()}  max_row={ws.max_row} max_col={ws.max_column}")
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=6, values_only=True), 1):
        cells = ["" if c is None else str(c) for c in row]
        print(f"r{i}: " + " | ".join(cells[:20]))
    # 最后两行
    rows = list(ws.iter_rows(min_row=max(1, ws.max_row - 1), max_row=ws.max_row, values_only=True))
    for j, row in enumerate(rows, ws.max_row - len(rows) + 1):
        cells = ["" if c is None else str(c) for c in row]
        print(f"r{j}: " + " | ".join(cells[:20]))
wb.close()
