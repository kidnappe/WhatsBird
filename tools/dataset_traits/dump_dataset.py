# -*- coding: utf-8 -*-
"""清点头文件行 + 统计实际物种行数；导出 CSV 供后续使用。"""
import csv
import sys
import openpyxl

src = sys.argv[1]
out_csv = sys.argv[2] if len(sys.argv) > 2 else None

wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
ws = wb['Chinesebirdsdata']
rows = list(ws.iter_rows(values_only=True))
h1, h2 = rows[0], rows[1]

print("列数:", len(h1))
for i in range(len(h1)):
    a = ("" if h1[i] is None else str(h1[i])).strip()
    b = ("" if i >= len(h2) or h2[i] is None else str(h2[i])).strip()
    print(f"  col{i+1:2d} ({openpyxl.utils.get_column_letter(i+1)}): 一级='{a}'  二级='{b}'")

data = [r for r in rows[2:] if any(c is not None and str(c).strip() != "" for c in r)]
latin_col = 4  # 种拉丁名
species = [r for r in data if r[latin_col] is not None and str(r[latin_col]).strip()]
print("\n非空行数:", len(data), " 有拉丁名行数:", len(species))
uniq = {str(r[latin_col]).strip() for r in species}
print("拉丁名去重:", len(uniq))
uniq_cn = {str(r[2]).strip() for r in species if r[2]}
print("中文名去重:", len(uniq_cn))

# 完整度速览：每列非空占比
print("\n列完整度（非空/物种数）:")
n = len(species)
for i in range(len(h1)):
    a = ("" if h1[i] is None else str(h1[i])).strip()
    b = ("" if i >= len(h2) or h2[i] is None else str(h2[i])).strip()
    label = f"{a}/{b}" if b else a
    cnt = sum(1 for r in species if i < len(r) and r[i] is not None and str(r[i]).strip() != "")
    print(f"  {label:<24s} {cnt:>5d}  {cnt/n*100:5.1f}%")

if out_csv:
    with open(out_csv, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow([f"{('' if h1[i] is None else h1[i])}|{('' if i>=len(h2) or h2[i] is None else h2[i])}".strip('|') for i in range(len(h1))])
        for r in species:
            w.writerow(["" if c is None else c for c in r])
    print("\n已写 CSV:", out_csv)
wb.close()
